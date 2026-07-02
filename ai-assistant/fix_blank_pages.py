#!/usr/bin/env python3
"""Remove empty Heading paragraphs that cause blank pages in front matter."""

from docx import Document
from pathlib import Path

DOC = Path("lisans_tez_60sayfa.docx")


def is_empty_heading(paragraph) -> bool:
    if paragraph.style is None or paragraph.style.name != "Heading":
        return False
    return not paragraph.text.strip()


def main() -> None:
    doc = Document(str(DOC))
    removed = 0

    # Walk backwards so indices stay valid when deleting
    for i in range(len(doc.paragraphs) - 1, -1, -1):
        p = doc.paragraphs[i]
        if is_empty_heading(p):
            el = p._element
            el.getparent().remove(el)
            removed += 1

    doc.save(str(DOC))
    print(f"Removed {removed} empty Heading paragraph(s) from {DOC}")


if __name__ == "__main__":
    main()
