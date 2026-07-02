#!/usr/bin/env python3
"""Rebuild thesis cover — compatible with Word Mac + Word Online."""
import copy
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

BASE = Path(__file__).parent
BACKUP = BASE / "05210000209_05230001155-Tez_BACKUP.docx"
OUT = BASE / "05210000209_05230001155-Tez.docx"
KAPAK_OUT = BASE / "kapakLisans.docx"
LOGO = BASE / "Ege_University_logo.svg.png"
BG_PNG = BASE / "cover_background.png"

W, H = 1240, 1754
SIDEBAR_W = 118
TEXT_COLOR = "000000"  # siyah — Online/Mac uyumluluğu
NAVY = "1F4E79"


def create_cover_background():
    img = Image.new("RGB", (W, H), "#FFFFFF")
    pixels = img.load()
    top = (255, 255, 255)
    bottom = (170, 198, 222)  # biraz daha koyu mavi alt
    for y in range(H):
        t = y / (H - 1)
        r = int(top[0] + (bottom[0] - top[0]) * t)
        g = int(top[1] + (bottom[1] - top[1]) * t)
        b = int(top[2] + (bottom[2] - top[2]) * t)
        for x in range(SIDEBAR_W, W):
            pixels[x, y] = (r, g, b)

    for x in range(SIDEBAR_W):
        t = x / max(SIDEBAR_W - 1, 1)
        shade = int(175 + 35 * (1 - abs(t - 0.3)))
        for y in range(H):
            pixels[x, y] = (shade, shade + 20, min(255, shade + 45))

    try:
        font = ImageFont.truetype(
            "/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf", 36
        )
    except OSError:
        font = ImageFont.load_default()

    text = "EÜ  MÜHENDİSLİK  FAKÜLTESİ"
    # dikey metin — ortada
    tb = Image.new("RGBA", (600, 80), (0, 0, 0, 0))
    tdraw = ImageDraw.Draw(tb)
    tdraw.text((0, 0), text, fill=(31, 78, 121, 255), font=font)
    tb = tb.rotate(90, expand=True)
    y_pos = (H - tb.height) // 2
    img.paste(tb, (12, y_pos), tb)
    img.save(BG_PNG, "PNG")
    print(f"Background: {BG_PNG}")


def nil_borders_element():
    tcB = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "nil")
        tcB.append(e)
    return tcB


def clear_table_borders(tbl):
    tblPr = tbl._tbl.tblPr
    shd = tblPr.find(qn("w:shd"))
    if shd is not None:
        tblPr.remove(shd)
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    tblB = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "nil")
        tblB.append(e)
    tblPr.append(tblB)
    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            for tag in (qn("w:shd"), qn("w:tcBorders")):
                el = tcPr.find(tag)
                if el is not None:
                    tcPr.remove(el)
            tcPr.append(nil_borders_element())


def set_run_color(run, hex_color=TEXT_COLOR):
    run.font.color.rgb = RGBColor.from_string(hex_color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("w14:textFill", "w:themeColor", "w:themeShade", "w:themeTint"):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    color = rPr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rPr.append(color)
    color.set(qn("w:val"), hex_color)


def tighten_paragraph(p, space_before=0, space_after=6, align_center=True):
    try:
        p.style = "Normal"
    except Exception:
        pass
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    if align_center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def style_cover_paragraph(p, bold=True, size_pt=14, space_after=6):
    tighten_paragraph(p, space_before=0, space_after=space_after)
    for run in p.runs:
        run.bold = bold
        run.font.name = "Times New Roman"
        run.font.size = Pt(size_pt)
        set_run_color(run)


def add_fullpage_background(paragraph, image_path, doc):
    part = doc.part
    r_id, _ = part.get_or_add_image(str(image_path))
    cx, cy = 7560000, 10692000
    xml = (
        f'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        f'<w:drawing>'
        f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="251658240" behindDoc="1" locked="1" layoutInCell="0" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/>'
        f'<wp:docPr id="99001" name="CoverBackground"/>'
        f'<wp:cNvGraphicFramePr/>'
        f'<a:graphic>'
        f'<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        f'<pic:pic>'
        f'<pic:nvPicPr><pic:cNvPr id="99001" name="CoverBackground"/>'
        f'<pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip r:embed="{r_id}"/>'
        f'<a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
        f'</pic:pic>'
        f'</a:graphicData>'
        f'</a:graphic>'
        f'</wp:anchor>'
        f'</w:drawing>'
        f"</w:r>"
    )
    paragraph._p.insert(0, parse_xml(xml))


def build_cover(doc):
    body = doc.element.body

    for p in doc.paragraphs[:20]:
        if p.text.strip() == "EÜ MÜHENDİSLİK FAKÜLTESİ":
            p._p.getparent().remove(p._p)
            break

    # Arka plan — ilk paragrafın EN BAŞINA (metinlerden önce, arkada)
    first = doc.paragraphs[0]
    bg_holder = first.insert_paragraph_before()
    tighten_paragraph(bg_holder, space_before=0, space_after=0, align_center=False)
    add_fullpage_background(bg_holder, BG_PNG, doc)

    # Logo
    logo_p = doc.paragraphs[0].insert_paragraph_before()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = logo_p.paragraph_format
    pf.left_indent = Cm(2.3)
    pf.space_before = Pt(2)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    try:
        logo_p.style = "Normal"
    except Exception:
        pass
    logo_p.add_run().add_picture(str(LOGO), width=Cm(2.5))

    # Kapak metinleri — sıkı aralık, siyah yazı
    spacing_map = {
        "EGE ÜNİVERSİTESİ": (16, 4),
        "LİSANS TEZİ": (14, 8),
    }
    title_started = False
    for p in doc.paragraphs[:22]:
        t = p.text.strip()
        if not t:
            tighten_paragraph(p, 0, 0)
            continue
        if t == "ÖZET":
            break
        if t in spacing_map:
            sz, sa = spacing_map[t]
            style_cover_paragraph(p, bold=True, size_pt=sz, space_after=sa)
        elif t.startswith("MCP TABANLI"):
            title_started = True
            style_cover_paragraph(p, bold=True, size_pt=13, space_after=2)
        elif title_started and t in (
            "GOOGLE VE NOTION ENTEGRASYONU İLE",
            "ÇOK KULLANICILI MİMARİ TASARIMI VE GELİŞTİRİLMESİ",
        ):
            style_cover_paragraph(p, bold=True, size_pt=13, space_after=2)
        elif t in ("Mustafa Yiğit GÜZEL", "Meriç ÖZKAYAĞAN"):
            style_cover_paragraph(p, bold=True, size_pt=14, space_after=2)
        elif t.startswith("Tez Danışmanı"):
            style_cover_paragraph(p, bold=True, size_pt=12, space_after=4)
        elif t == "Bilgisayar Mühendisliği Anabilim Dalı":
            style_cover_paragraph(p, bold=False, size_pt=12, space_after=4)
        elif t.startswith("Sunuş Tarihi"):
            style_cover_paragraph(p, bold=True, size_pt=12, space_after=4)
        elif t in ("Bornova-İZMİR", "2026"):
            style_cover_paragraph(p, bold=True, size_pt=12, space_after=4)

    # Section break
    cover_break_p = None
    for i, p in enumerate(doc.paragraphs[:25]):
        if p.text.strip() == "ÖZET":
            cover_break_p = doc.paragraphs[i - 1]
            break

    body_sectPr = body.find(qn("w:sectPr"))
    cover_sectPr = copy.deepcopy(body_sectPr)
    for ref in cover_sectPr.findall(qn("w:headerReference")) + cover_sectPr.findall(
        qn("w:footerReference")
    ):
        cover_sectPr.remove(ref)
    for pgB in cover_sectPr.findall(qn("w:pgBorders")):
        cover_sectPr.remove(pgB)

    pPr = cover_break_p._p.get_or_add_pPr()
    old = pPr.find(qn("w:sectPr"))
    if old is not None:
        pPr.remove(old)
    pPr.append(cover_sectPr)


def export_kapak_only(src_path, dst_path):
    shutil.copy2(src_path, dst_path)
    d = Document(dst_path)
    body = d.element.body
    ozet_el = None
    for p in d.paragraphs:
        if p.text.strip() == "ÖZET":
            ozet_el = p._p
            break
    if ozet_el is None:
        return
    removing = False
    for child in list(body.iterchildren()):
        if child is ozet_el:
            removing = True
        if removing:
            body.remove(child)
    for sp in body.findall(".//" + qn("w:sectPr")):
        for ref in sp.findall(qn("w:headerReference")) + sp.findall(qn("w:footerReference")):
            sp.remove(ref)
    d.save(dst_path)


def main():
    create_cover_background()
    shutil.copy2(BACKUP, OUT)
    doc = Document(str(OUT))
    for ti in (0, 1, 2):
        clear_table_borders(doc.tables[ti])
    build_cover(doc)
    doc.save(str(OUT))
    export_kapak_only(OUT, KAPAK_OUT)
    print(f"Saved: {OUT}")
    print(f"Saved: {KAPAK_OUT}")


if __name__ == "__main__":
    main()
